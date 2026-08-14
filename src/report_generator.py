"""Generate a traceable Word analysis report from the current detection result."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from typing import Any, Mapping, Optional, Union

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .business_metrics import get_severity_count


PRIMARY = "1A3C5E"
ACCENT = "E8412B"
MUTED = "667085"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def _rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(*(int(value[index : index + 2], 16) for index in (0, 2, 4)))


def _set_run_font(run, *, size: Optional[float] = None, bold: Optional[bool] = None,
                  color: Optional[str] = None, name: str = "Calibri") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = _rgb(color)


def _set_cell_background(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color.lstrip("#"))


def _set_cell_margins(cell, top: int = 80, start: int = 120,
                      bottom: int = 80, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError("Word 表格列宽总和必须等于 9360 DXA。")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr

    table_width = properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.first_child_found_in("w:tcW")
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, PRIMARY, 16, 8),
        "Heading 2": (13, PRIMARY, 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("Shipping Anomaly Detection  |  Analysis Report")
    _set_run_font(header_run, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("Portfolio Final Edition   |   ")
    _set_run_font(label, size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "Shipping Anomaly Detection Report"


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def _add_kpi_table(doc: Document, stats: dict) -> None:
    kpis = [
        ("总订单量", f"{stats['total_orders']:,}"),
        ("异常订单", f"{stats['total_anomalies']:,}"),
        ("异常率", f"{stats['anomaly_rate'] * 100:.1f}%"),
        ("平均延误", f"{stats['avg_delay']:.1f} 天"),
        ("严重异常", f"{get_severity_count(stats, '严重'):,}"),
    ]
    table = doc.add_table(rows=1, cols=len(kpis))
    table.style = "Table Grid"
    widths = [1872] * 5
    for index, (label, value) in enumerate(kpis):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_background(cell, PRIMARY)
        value_paragraph = cell.paragraphs[0]
        value_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        value_paragraph.paragraph_format.space_after = Pt(2)
        value_run = value_paragraph.add_run(value)
        _set_run_font(value_run, size=15, bold=True, color=WHITE)
        label_paragraph = cell.add_paragraph()
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_paragraph.paragraph_format.space_after = Pt(0)
        label_run = label_paragraph.add_run(label)
        _set_run_font(label_run, size=8, color="D6E2EE")
    _set_table_geometry(table, widths)


def _value_text(value: Any) -> str:
    if pd.isna(value):
        return "—"
    if isinstance(value, float):
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def _add_data_table(
    doc: Document,
    frame: pd.DataFrame,
    widths: list[int],
    *,
    percent_columns: Optional[set[str]] = None,
) -> None:
    percent_columns = percent_columns or set()
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    _repeat_header(table.rows[0])
    for index, column in enumerate(frame.columns):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_background(cell, PRIMARY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(column))
        _set_run_font(run, size=8.5, bold=True, color=WHITE)

    for row_index, (_, row_data) in enumerate(frame.iterrows()):
        row = table.add_row()
        for column_index, column in enumerate(frame.columns):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 0:
                _set_cell_background(cell, "F8F9FA")
            value = row_data[column]
            if column in percent_columns and not pd.isna(value):
                text = f"{float(value) * 100:.1f}%"
            else:
                text = _value_text(value)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column in {"承运商", "路线", "样本提示"} else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            _set_run_font(run, size=8.5)
    _set_table_geometry(table, widths)


def _add_chart(doc: Document, chart_images: Mapping[str, Any], key: str) -> bool:
    if key not in chart_images:
        return False
    source = chart_images[key]
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(bytes(source))
    doc.add_picture(source, width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
    return True


def _markdown_lines(text: str) -> list[str]:
    lines = []
    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        line = re.sub(r"^#{1,6}\s*", "", line)
        line = re.sub(r"^[-*>]\s*", "", line)
        line = re.sub(r"\*\*([^*]+)\*\*", r"\1", line)
        line = line.replace("`", "")
        lines.append(line)
    return lines


def _recommendations(stats: dict) -> list[str]:
    recommendations = []
    severe = get_severity_count(stats, "严重")
    if severe:
        recommendations.append(f"优先人工复核 {severe} 条严重异常，并补充港口、天气、船期或操作记录确认根因。")
    carriers = stats.get("carrier_stats")
    if isinstance(carriers, pd.DataFrame) and not carriers.empty:
        top = carriers.iloc[0]
        note = "（样本较少，先补充观察）" if top["样本提示"] == "样本较少" else ""
        recommendations.append(
            f"对 {top['承运商']} 联合检查订单量 {int(top['总订单数'])}、异常数 "
            f"{int(top['异常订单数'])} 和异常率 {float(top['异常率']) * 100:.1f}%{note}。"
        )
    routes = stats.get("route_stats")
    if isinstance(routes, pd.DataFrame) and not routes.empty:
        top = routes.iloc[0]
        recommendations.append(
            f"复核路线 {top['路线']} 的计划运输时长与节点衔接，并持续按月观察调整后风险排名。"
        )
    recommendations.append("将异常原因复核结果回填到运营台账，定期校准业务阈值，不以单次模型结果替代业务判断。")
    return recommendations


def generate_report(
    stats: dict,
    result_df: pd.DataFrame,
    chart_images: Mapping[str, Any],
    output_path: Optional[Union[str, Path]] = None,
    *,
    insight: Optional[str] = None,
    insight_label: str = "自动业务摘要",
    throughput_available: bool = False,
) -> bytes:
    """Create the report, optionally save it, and always return DOCX bytes."""

    doc = Document()
    _configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(42)
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run("航运运输异常检测分析报告")
    _set_run_font(title_run, size=26, bold=True, color=PRIMARY)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    subtitle_run = subtitle.add_run("Shipping Anomaly Detection  |  Final Portfolio Edition")
    _set_run_font(subtitle_run, size=11, color=MUTED)

    _add_heading(doc, "一、执行摘要")
    _add_kpi_table(doc, stats)
    doc.add_paragraph(
        f"本次分析覆盖 {stats['total_orders']:,} 条订单，识别异常 {stats['total_anomalies']:,} 条，"
        f"异常率 {stats['anomaly_rate'] * 100:.1f}%，其中严重异常 "
        f"{get_severity_count(stats, '严重')} 条。所有结果均由当前输入数据实时计算，"
        "综合业务规则、路线分组 Z-Score 与 Isolation Forest，并保留逐条解释。"
    )

    _add_heading(doc, "二、数据概览")
    date_range = stats.get("date_range", {})
    doc.add_paragraph(
        f"数据时间范围：{date_range.get('start') or '未知'} 至 {date_range.get('end') or '未知'}；"
        f"平均计划运输时长 {stats['avg_planned_transit']:.1f} 天；"
        f"平均延误 {stats['avg_delay']:.1f} 天；最长延误 {stats['max_delay']:.1f} 天。"
    )
    quality = stats.get("data_quality", {})
    if quality:
        doc.add_paragraph(
            f"上传记录 {quality.get('raw_rows', stats['total_orders'])} 条，"
            f"可用记录 {quality.get('valid_rows', stats['total_orders'])} 条，"
            f"清洗剔除 {quality.get('dropped_rows', 0)} 条。"
        )

    _add_heading(doc, "三、异常检测结果")
    _add_chart(doc, chart_images, "anomaly_overview")
    _add_chart(doc, chart_images, "delay_distribution")
    method_counts = stats.get("method_counts", {})
    doc.add_paragraph(
        f"业务规则命中 {method_counts.get('业务规则', 0)} 条，分组 Z-Score 命中 "
        f"{method_counts.get('分组 Z-Score', 0)} 条，Isolation Forest 命中 "
        f"{method_counts.get('Isolation Forest', 0)} 条。业务规则可直接触发异常；"
        "否则需至少两种独立方法一致命中。"
    )

    _add_heading(doc, "四、严重等级分布")
    severity_frame = pd.DataFrame(
        {
            "严重等级": ["严重", "中等", "轻微", "正常"],
            "订单数": [get_severity_count(stats, level) for level in ("严重", "中等", "轻微", "正常")],
        }
    )
    severity_frame["占比"] = severity_frame["订单数"] / stats["total_orders"]
    _add_data_table(doc, severity_frame, [2400, 2400, 4560], percent_columns={"占比"})

    _add_heading(doc, "五、承运商分析")
    _add_chart(doc, chart_images, "carrier_performance")
    carrier_frame = stats["carrier_stats"].head(10)[
        ["承运商", "总订单数", "异常订单数", "异常率", "平均延误天数", "样本提示"]
    ]
    _add_data_table(
        doc,
        carrier_frame,
        [1700, 1200, 1300, 1300, 1700, 2160],
        percent_columns={"异常率"},
    )

    _add_heading(doc, "六、路线风险")
    if throughput_available and _add_chart(doc, chart_images, "port_throughput"):
        doc.add_paragraph("港口吞吐量仅用于补充运营背景，不参与订单异常模型和投票。")
    else:
        doc.add_paragraph("本次未提供港口吞吐量数据，因此不生成港口运营背景图；核心订单检测不受影响。")
    route_frame = stats["route_stats"].head(10)[
        ["路线", "订单数", "异常数", "异常率", "平均延误", "样本提示"]
    ]
    _add_data_table(
        doc,
        route_frame,
        [2600, 1100, 1100, 1200, 1400, 1960],
        percent_columns={"异常率"},
    )

    _add_heading(doc, "七、重点异常订单")
    top = stats.get("top_anomalies", pd.DataFrame())
    if isinstance(top, pd.DataFrame) and not top.empty:
        for _, row in top.iterrows():
            heading = doc.add_paragraph(style="Heading 2")
            heading.add_run(
                f"{row['订单ID']}  |  {row['severity']}  |  风险分数 {row['anomaly_score']}"
            )
            doc.add_paragraph(
                f"承运商：{row['承运商']}；路线：{row['路线']}；"
                f"命中方法：{row['anomaly_methods']}；实际延误：{row['实际延误天数']} 天。"
            )
            doc.add_paragraph(f"异常原因：{row['reason']}")
    else:
        doc.add_paragraph("当前数据没有达到综合异常条件的订单。")

    _add_heading(doc, "八、AI / 自动业务洞察")
    label_paragraph = doc.add_paragraph()
    label_run = label_paragraph.add_run(insight_label)
    _set_run_font(label_run, size=11, bold=True, color=PRIMARY)
    lines = _markdown_lines(insight or "本次未生成业务摘要。")
    for line in lines:
        doc.add_paragraph(line)

    _add_heading(doc, "九、建议")
    for index, recommendation in enumerate(_recommendations(stats), start=1):
        paragraph = doc.add_paragraph()
        lead = paragraph.add_run(f"建议 {index}｜")
        _set_run_font(lead, bold=True, color=PRIMARY)
        paragraph.add_run(recommendation)

    _add_heading(doc, "十、方法说明")
    config = stats.get("detector_config", {})
    doc.add_paragraph(
        f"业务规则用于识别明确的延误或运费数据风险；分组 Z-Score 以路线为基线，"
        f"当前阈值为 {config.get('zscore_threshold', '未记录')}，小样本路线不参与统计判断；"
        f"Isolation Forest contamination 为 {config.get('iforest_contamination', '未记录')}，"
        "仅使用经校验的数值特征并固定随机种子。严重等级由最终异常状态、命中方法数和关键延误规则集中计算。"
    )
    note = doc.add_paragraph(
        "免责声明：示例数据为模拟数据。本平台用于数据分析、研究及项目展示，不代表真实航运企业生产系统；"
        "模型结果应结合业务记录进行人工复核。"
    )
    note_run = note.runs[0]
    _set_run_font(note_run, size=8.5, color=MUTED)

    buffer = BytesIO()
    doc.save(buffer)
    content = buffer.getvalue()
    if output_path is not None:
        Path(output_path).write_bytes(content)
    return content
