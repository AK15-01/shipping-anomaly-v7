from __future__ import annotations

from io import BytesIO

from docx import Document

from src.ai_analyst import generate_local_summary
from src.business_metrics import build_summary_stats
from src.report_generator import generate_report


def _all_document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def test_report_generation_contains_all_major_sections(metrics_result):
    stats = build_summary_stats(metrics_result)
    stats["detector_config"] = {
        "zscore_threshold": 2.5,
        "iforest_contamination": 0.05,
    }
    content = generate_report(
        stats,
        metrics_result,
        {},
        insight=generate_local_summary(stats),
        insight_label="自动业务摘要",
    )
    assert content[:2] == b"PK"
    document = Document(BytesIO(content))
    text = _all_document_text(document)
    for heading in [
        "一、执行摘要",
        "二、数据概览",
        "三、异常检测结果",
        "四、严重等级分布",
        "五、承运商分析",
        "六、路线风险",
        "七、重点异常订单",
        "八、AI / 自动业务洞察",
        "九、建议",
        "十、方法说明",
    ]:
        assert heading in text
    assert "自动业务摘要" in text


def test_report_uses_same_severe_count_as_statistics(metrics_result):
    stats = build_summary_stats(metrics_result)
    content = generate_report(stats, metrics_result, {}, insight="自动业务摘要")
    text = _all_document_text(Document(BytesIO(content)))
    assert "严重异常 9 条" in text
